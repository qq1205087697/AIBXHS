"""G1/G2/G6 缺口补齐的端到端验证（可重复执行，跑完自清理）

覆盖：
- G1 tag_locked：search_inventory 返回体 + get_summary_children 返回体
- G2 filter_status：get_summary_children 返回体 + export_inventory_to_excel 列/中文映射
- G6 F10 兜底报告导入：services/startop_report_importer.import_report_file
      · 正常导入（列名映射 + 跳过不可归因行）
      · 文件哈希幂等（二次上传 DUPLICATE）
      · 非法文件 FAILED + 流水留痕

安全：所有测试数据使用哨兵值（sku='__g1_test__' / store_id='__G6_TEST__'），
      finally 中按哨兵精确删除，不触碰既有业务数据。
"""
import io
import os
import sys
import traceback
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pandas as pd  # noqa: E402
from sqlalchemy import text  # noqa: E402

from database.database import SessionLocal  # noqa: E402
from services.tag_service import _ensure_models_loaded  # noqa: E402

TENANT = 1
SKU_SENTINEL = "__g1_test__"
STORE_SENTINEL = "__G6_TEST__"

results = []


def check(name, ok, detail=""):
    results.append((name, bool(ok), detail))
    print(f"[{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail else ""))


def build_report_bytes(with_bad_row=True):
    rows = [
        {"搜索词": "halloween decorations", "ASIN": "B0G6TEST01", "数据日期": date.today().isoformat(),
         "站点": "US", "店铺": STORE_SENTINEL, "词类型": "出单词"},
        {"搜索词": "pumpkin lights", "ASIN": "B0G6TEST01", "数据日期": date.today().isoformat(),
         "站点": "US", "店铺": STORE_SENTINEL, "词类型": "出单词"},
        {"搜索词": "spooky garland", "ASIN": "B0G6TEST02", "数据日期": date.today().isoformat(),
         "站点": "US", "店铺": STORE_SENTINEL, "词类型": "流量词"},
    ]
    if with_bad_row:
        # 缺 ASIN → 不可归因，应计入 rows_skipped
        rows.append({"搜索词": "no asin here", "ASIN": None, "数据日期": date.today().isoformat(),
                     "站点": "US", "店铺": STORE_SENTINEL, "词类型": "出单词"})
    buf = io.BytesIO()
    pd.DataFrame(rows).to_excel(buf, index=False, engine="openpyxl")
    return buf.getvalue()


def main():
    _ensure_models_loaded()
    db = SessionLocal()

    from services.inventory_service import (
        search_inventory, get_summary_children, export_inventory_to_excel,
    )
    from services.startop_report_importer import import_report_file

    # ============ 基线：找一条当前可见（KEEP）的 ASIN 作对照组 ============
    baseline = search_inventory(db, tenant_id=TENANT, page=1, page_size=5)
    check("基线 search_inventory 可查询", baseline["total"] > 0,
          f"total={baseline['total']}")
    if not baseline["items"]:
        raise SystemExit("基线列表为空，无法继续 G1 验证")
    # 对照 ASIN 必须「完全没有标签行」：自 2026-09-29 起标签只做「人工抑制」（active=0 但行保留），
    # 若对照 ASIN 已有同键历史行，下面的哨兵 INSERT 会撞 uk_tag 唯一键。
    _tagged = {r[0] for r in db.execute(text(
        "SELECT DISTINCT asin FROM product_tags WHERE tenant_id=:t"), {"t": TENANT}).all()}
    _pick = next((it for it in baseline["items"] if it["asin"] not in _tagged), None)
    control_asin = (_pick or baseline["items"][0])["asin"]
    check("基线返回体已含 tag_locked 字段且默认 0",
          "tag_locked" in baseline["items"][0] and baseline["items"][0]["tag_locked"] == 0,
          f"control asin={control_asin} value={baseline['items'][0].get('tag_locked')}")

    # ============ G1：人工锁定标识（临时插入哨兵标签） ============
    db.execute(text("""
        INSERT INTO product_tags
        (tenant_id, asin, sku, tag_type, tag_value, source, locked, active, reason, created_at, updated_at)
        VALUES (:tid, :asin, :sku, 'force_replenish', '1', 'manual', 1, 1, 'G1 verify temp', NOW(), NOW())
    """), {"tid": TENANT, "asin": control_asin, "sku": SKU_SENTINEL})
    db.commit()

    locked_view = search_inventory(db, tenant_id=TENANT, keyword=control_asin, page=1, page_size=20)
    hit = next((it for it in locked_view["items"] if it["asin"] == control_asin), None)
    check("G1 search_inventory.tag_locked 随人工锁定标签置 1",
          hit is not None and hit.get("tag_locked") == 1,
          f"asin={control_asin} tag_locked={None if hit is None else hit.get('tag_locked')}")

    # 回滚哨兵标签 → tag_locked 复位 0
    db.execute(text("DELETE FROM product_tags WHERE sku = :sku"), {"sku": SKU_SENTINEL})
    db.commit()
    after = search_inventory(db, tenant_id=TENANT, keyword=control_asin, page=1, page_size=20)
    hit2 = next((it for it in after["items"] if it["asin"] == control_asin), None)
    check("G1 解除锁定后 tag_locked 复位 0",
          hit2 is not None and hit2.get("tag_locked") == 0,
          f"tag_locked={None if hit2 is None else hit2.get('tag_locked')}")

    # ============ G1/G2：子行返回体 ============
    # 必须挑「非节日且未停售」的汇总行：get_summary_children 按 K3 只返回 KEEP 子行，
    # 若父行本身是节日商品，子行会被正确排除（children=0）→ 断言会误判为失败。
    parent = db.execute(text("""
        SELECT asin, account FROM inventory_snapshots
        WHERE tenant_id=:tid AND summary_flag='是' AND account IS NOT NULL
          AND (is_holiday = '' OR is_holiday IS NULL)
          AND (is_discontinued = 0 OR is_discontinued IS NULL)
          AND snapshot_date=(SELECT MAX(snapshot_date) FROM inventory_snapshots WHERE tenant_id=:tid)
        LIMIT 1
    """), {"tid": TENANT}).first()
    if parent:
        children = get_summary_children(db, TENANT, parent[0], parent[1])
        check("G1/G2 子行返回体已含 filter_status / tag_locked / force_replenish",
              all(k in c for c in children for k in ("filter_status", "tag_locked", "force_replenish"))
              if children else False,
              f"parent={parent[0]} children={len(children)}")
        check("G2 子行 filter_status 全部为 KEEP（与主列表同口径）",
              all(c["filter_status"] == "KEEP" for c in children) if children else False,
              f"statuses={sorted({c['filter_status'] for c in children}) if children else []}")
        check("G1 子行 tag_locked 为 0/1 整数",
              all(c["tag_locked"] in (0, 1) for c in children) if children else False)
    else:
        check("G1/G2 子行返回体（未找到父汇总行，跳过）", False, "无 summary_flag='是' 数据")

    # ============ G2：导出列与中文映射 ============
    items = search_inventory(db, tenant_id=TENANT, page=1, page_size=5)["items"]
    xlsx = export_inventory_to_excel(items, ["tag_locked", "filter_status", "force_replenish"])
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(xlsx.getvalue())).active
    headers = [c.value for c in ws[1]]
    check("G2 导出表头含 人工锁定 / 过滤状态 / 强制补货",
          headers[:3] == ["人工锁定", "过滤状态", "强制补货"], f"headers={headers}")
    col_values = {ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)}
    check("G2 filter_status 导出为中文枚举（非 KEEP 原文）",
          col_values <= {"参与补货", "停售", "负利润排除", "节日排除"}, f"values={col_values}")

    # ============ G6：F10 兜底报告导入 ============
    content = build_report_bytes()
    r1 = import_report_file(db, TENANT, "g6_test_report.xlsx", content)
    check("G6 报告导入 SUCCESS", r1["status"] == "SUCCESS", str(r1["message"]))
    check("G6 行数统计正确（3 入库 / 1 跳过）",
          r1["rows_total"] == 4 and r1["rows_imported"] == 3 and r1["rows_skipped"] == 1,
          f"total={r1['rows_total']} imported={r1['rows_imported']} skipped={r1['rows_skipped']}")

    kinds = dict(db.execute(text(
        "SELECT word_type, COUNT(*) FROM startop_word_data WHERE store_id=:s GROUP BY word_type"
    ), {"s": STORE_SENTINEL}).all())
    check("G6 词类型映射正确（出单词 2 / 流量词 1）",
          kinds.get("search_term") == 2 and kinds.get("traffic_term") == 1, f"kinds={kinds}")

    src = db.execute(text(
        "SELECT DISTINCT data_source FROM startop_word_data WHERE store_id=:s"
    ), {"s": STORE_SENTINEL}).scalars().all()
    check("G6 落库 data_source='report'", src == ["report"], f"data_source={src}")

    r2 = import_report_file(db, TENANT, "g6_test_report.xlsx", content)
    check("G6 文件哈希幂等（二次上传返回 DUPLICATE 且不写库）",
          r2["status"] == "DUPLICATE" and r2["rows_imported"] is None, str(r2["message"]))
    cnt = db.execute(text(
        "SELECT COUNT(*) FROM startop_word_data WHERE store_id=:s"
    ), {"s": STORE_SENTINEL}).scalar()
    check("G6 幂等后行数未增长", cnt == 3, f"rows={cnt}")

    r3 = import_report_file(db, TENANT, "g6_bad_file.xlsx", b"this is not an excel file")
    check("G6 非法文件返回 FAILED", r3["status"] == "FAILED", str(r3["message"])[:80])
    log = db.execute(text(
        "SELECT status, fallback_ready FROM startop_pull_logs WHERE target_scope LIKE 'report:g6_bad_file%'"
    )).first()
    check("G6 失败导入写流水（FAILED + fallback_ready=1）",
          log is not None and tuple(log) == ("FAILED", 1), f"log={tuple(log) if log else None}")
    return results


def cleanup():
    db = SessionLocal()
    try:
        db.execute(text("DELETE FROM product_tags WHERE sku = :sku"), {"sku": SKU_SENTINEL})
        db.execute(text("DELETE FROM startop_word_data WHERE store_id = :s"), {"s": STORE_SENTINEL})
        db.execute(text("DELETE FROM startop_pull_logs WHERE target_scope LIKE 'report:g6_%'"))
        db.commit()
        left_snaps = db.execute(text("SELECT COUNT(*) FROM startop_word_data WHERE store_id=:s"),
                                {"s": STORE_SENTINEL}).scalar()
        left_tags = db.execute(text("SELECT COUNT(*) FROM product_tags WHERE sku=:sku"),
                               {"sku": SKU_SENTINEL}).scalar()
        left_logs = db.execute(text("SELECT COUNT(*) FROM startop_pull_logs WHERE target_scope LIKE 'report:g6_%'")).scalar()
        print(f"[CLEANUP] 残留：word_data={left_snaps} tags={left_tags} logs={left_logs}")
    finally:
        db.close()


if __name__ == "__main__":
    try:
        main()
    except Exception:
        traceback.print_exc()
        results.append(("脚本异常", False, "见 traceback"))
    finally:
        cleanup()
        total = len(results)
        passed = sum(1 for _, ok, _ in results if ok)
        print(f"\n===== {passed}/{total} PASS =====")
        sys.exit(0 if passed == total else 1)
