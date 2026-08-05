from fastapi import APIRouter, Depends, Query, UploadFile, File, Request
from sqlalchemy.orm import Session
from sqlalchemy import text
from typing import Optional, List
from pydantic import BaseModel
from database.database import get_db
from models.product_page_info import ProductPageInfo
from models.store import Store
from models.department import UserDepartment
from dependencies import get_current_user
from models.user import User
import httpx
import asyncio
import openpyxl
import io
import logging

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/product-page-info", tags=["product-page-info"])


def is_admin_user(user: User, db: Session) -> bool:
    """判断用户是否是管理员（通过 role_id）"""
    if not user.role_id:
        return False
    role = db.execute(text("""
        SELECT code FROM roles WHERE id = :role_id AND deleted_at IS NULL
    """), {"role_id": user.role_id}).fetchone()
    return role and role[0] == "admin"


def get_user_store_names(user: User, db: Session) -> list:
    """获取非管理员用户所属部门下的所有店铺inventory_name列表"""
    dept_ids = db.query(UserDepartment.department_id).filter(
        UserDepartment.user_id == user.id
    ).all()
    dept_id_list = [d[0] for d in dept_ids]
    if not dept_id_list:
        return []
    stores = db.query(Store.inventory_name).filter(
        Store.tenant_id == user.tenant_id,
        Store.department_id.in_(dept_id_list),
        Store.inventory_name.isnot(None)
    ).all()
    return [s[0] for s in stores if s[0]]


def parse_string_score(val: Optional[str]) -> Optional[float]:
    """将字符串评分（如"4.5"）转换为浮点数"""
    if not val:
        return None
    try:
        return float(val)
    except ValueError:
        return None

def calc_star_rating_score(val: Optional[float]) -> Optional[int]:
    """将星级评分转换为分数：
    0.1-3.9 → 0分, 4-4.9 → 8分, 0 → 6分, 5 → 10分
    """
    if val is None:
        return None
    if val == 0:
        return 6
    if val == 5:
        return 10
    if 4 <= val < 5:
        return 8
    if 0.1 <= val <= 3.9:
        return 0
    return None

def calc_image_count_score(count: int) -> int:
    """图片数评分：≥9 → 10分, ≥6 → 7分, ≥3 → 4分, <3 → 0分"""
    if count >= 9:
        return 10
    elif count >= 6:
        return 7
    elif count >= 3:
        return 4
    else:
        return 0

def calc_competitor_price_score(current_price_str: str, competitor_str: Optional[str]) -> tuple[int, int]:
    """竞品价格评分：
    competitor_price 格式：一行一个，每项形如 "ASIN:价格" 或 "ASIN：价格"
    将所有竞品价格与当前价格比较，按排名给分（满分15）：
    - 最便宜 → 15分
    - 第二便宜 → 12分
    - 第三便宜 → 8分
    - 第四便宜 → 4分
    - 第五名及以后 → 0分
    返回 (rank, score)
    """
    if not competitor_str or competitor_str.strip() == "":
        return (0, 0)

    try:
        current_price = float(current_price_str.replace('$', '').replace(',', '').strip())
    except ValueError:
        return (0, 0)

    prices = []
    for line in competitor_str.strip().splitlines():
        line = line.strip()
        if not line:
            continue
        # 支持 : 和 ：两种分隔符
        for sep in [':', '：']:
            if sep in line:
                _, price_part = line.split(sep, 1)
                try:
                    price = float(price_part.replace('$', '').replace(',', '').strip())
                    prices.append(price)
                except ValueError:
                    continue
                break

    if not prices:
        return (0, 0)

    all_prices = prices + [current_price]
    all_prices.sort()
    rank = all_prices.index(current_price) + 1

    if rank == 1:
        return (rank, 15)
    elif rank == 2:
        return (rank, 12)
    elif rank == 3:
        return (rank, 8)
    elif rank == 4:
        return (rank, 4)
    else:
        return (rank, 0)

def calc_total_score(
    title_rating: Optional[str],
    description_rating: Optional[str],
    keywords_rating: Optional[str],
    image_rating: Optional[str],
    star_rating_score: Optional[int],
    has_ad: Optional[int],
    has_aplus: int,
    has_video: bool,
    competitor_price_score: int = 0,
) -> int:
    """计算总分（满分100）

    各项分值：
    - 标题评分：15分
    - 描述评分：10分
    - 关键词评分：10分
    - 图片评分：15分
    - 星级评分：10分
    - has_ad：10分（>=3=10，2=8，1=6，0/空=0）
    - has_aplus：10分（0=0，1=5，2=10）
    - has_video：5分（0=0，1=5）
    - 竞品价格：15分（排名1=15，2=12，3=8，4=4，5+=0）
    """
    total = 0

    # 标题评分（满分15）
    ts = parse_string_score(title_rating)
    if ts is not None:
        total += min(ts, 15)

    # 描述评分（满分10）
    ds = parse_string_score(description_rating)
    if ds is not None:
        total += min(ds, 10)

    # 关键词评分（满分10）
    ks = parse_string_score(keywords_rating)
    if ks is not None:
        total += min(ks, 10)

    # 图片评分（满分15）
    is_ = parse_string_score(image_rating)
    if is_ is not None:
        total += min(is_, 15)

    # 星级评分（满分10）
    if star_rating_score is not None:
        total += star_rating_score

    # has_ad（满分10，按广告数量计分）
    if has_ad is not None and has_ad >= 3:
        total += 10
    elif has_ad is not None and has_ad == 2:
        total += 8
    elif has_ad is not None and has_ad == 1:
        total += 6
    else:
        total += 0

    # has_aplus（满分10）
    if has_aplus == 2:
        total += 10
    elif has_aplus == 1:
        total += 5
    else:
        total += 0

    # has_video（满分5）
    total += 5 if has_video else 0

    # 竞品价格（满分15）
    total += competitor_price_score

    return min(total, 100)


class ProductPageInfoResponse(BaseModel):
    id: int
    tenant_id: int
    asin: Optional[str] = None
    sku: Optional[str] = None
    store: Optional[str] = None
    store_original: Optional[str] = None
    title: Optional[str] = None
    keywords: Optional[str] = None
    product_description: Optional[str] = None
    bullet_points: Optional[str] = None
    price: Optional[str] = None
    image_count: Optional[int] = None
    title_rating: Optional[str] = None
    description_rating: Optional[str] = None
    keywords_rating: Optional[str] = None
    image_rating: Optional[str] = None
    star_rating: Optional[float] = None
    star_rating_score: Optional[int] = None
    competitor_price: Optional[str] = None
    competitor_price_rank: int = 0
    competitor_price_score: int = 0
    has_ad: Optional[bool] = False
    has_aplus: Optional[int] = 0
    has_video: Optional[bool] = False
    rating_status: int = 0
    total_score: int = 0
    traffic_keywords: Optional[str] = None

    class Config:
        from_attributes = True


@router.get("/")
async def get_product_page_info_list(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    asin_search: Optional[str] = Query(None),
    sku_search: Optional[str] = Query(None),
    store_filter: Optional[str] = Query(None),
    rating_status: Optional[int] = Query(None, description="评分状态筛选：0=未评分，1=已评分"),
    low_score: Optional[bool] = Query(None, description="低分筛选：已评分且总分<60"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """获取产品页面信息列表"""
    query = db.query(ProductPageInfo).filter(
        ProductPageInfo.tenant_id == current_user.tenant_id
    )

    # 非管理员用户按部门过滤数据
    if not is_admin_user(current_user, db):
        user_store_names = get_user_store_names(current_user, db)
        if user_store_names:
            query = query.filter(ProductPageInfo.store.in_(user_store_names))
        else:
            # 用户没有分配任何部门店铺，不显示任何数据
            query = query.filter(ProductPageInfo.id == -1)

    if rating_status is not None:
        query = query.filter(ProductPageInfo.rating_status == rating_status)

    if asin_search:
        query = query.filter(ProductPageInfo.asin.ilike(f"%{asin_search}%"))
    if sku_search:
        query = query.filter(ProductPageInfo.sku.ilike(f"%{sku_search}%"))
    if store_filter:
        # 先查找该 shop_abbr 对应的所有 inventory_name
        matching_stores = db.query(Store.inventory_name).filter(
            Store.tenant_id == current_user.tenant_id,
            Store.shop_abbr == store_filter
        ).all()
        matching_names = [s[0] for s in matching_stores if s[0]]
        if matching_names:
            query = query.filter(ProductPageInfo.store.in_(matching_names))
        else:
            # 没有匹配的店铺，回退到模糊搜索
            query = query.filter(ProductPageInfo.store.ilike(f"%{store_filter}%"))

    if low_score:
        # 低分筛选：先查所有已评分记录，计算总分后过滤
        all_rated = query.filter(ProductPageInfo.rating_status == 1).all()

        # 查询stores表映射
        store_rows = db.query(Store.inventory_name, Store.shop_abbr, Store.site).filter(
            Store.tenant_id == current_user.tenant_id,
            Store.inventory_name.isnot(None)
        ).all()
        store_abbr_map = {}
        for r in store_rows:
            if r.inventory_name:
                parts = [p for p in [r.shop_abbr, r.site] if p]
                store_abbr_map[r.inventory_name] = "-".join(parts) if parts else r.inventory_name

        filtered = []
        for item in all_rated:
            has_aplus_val = 0
            if isinstance(item.has_aplus, int):
                has_aplus_val = item.has_aplus
            elif isinstance(item.has_aplus, bool):
                has_aplus_val = 1 if item.has_aplus else 0
            has_ad_val = int(item.has_ad) if item.has_ad is not None else None
            has_video_val = bool(item.has_video) if item.has_video is not None else False
            star_score = calc_star_rating_score(item.star_rating)
            _, competitor_score = calc_competitor_price_score(item.price, item.competitor_price)
            total_score = calc_total_score(
                item.title_rating, item.description_rating, item.keywords_rating, item.image_rating,
                star_score, has_ad_val, has_aplus_val, has_video_val, competitor_score,
            )
            if total_score < 60:
                display_store = store_abbr_map.get(item.store, item.store) if item.store else item.store
                filtered.append((item, total_score, display_store, has_ad_val, has_aplus_val, has_video_val, star_score, competitor_score))

        # 按updated_at降序排序
        filtered.sort(key=lambda x: x[0].updated_at if x[0].updated_at else None, reverse=True)

        total = len(filtered)
        paged = filtered[(page - 1) * page_size : page * page_size]

        result = []
        for item, total_score, display_store, has_ad_val, has_aplus_val, has_video_val, star_score, competitor_score in paged:
            competitor_rank, competitor_score2 = calc_competitor_price_score(item.price, item.competitor_price)
            result.append({
                "id": item.id,
                "tenant_id": item.tenant_id,
                "asin": item.asin,
                "sku": item.sku,
                "store": display_store,
                "store_original": item.store,
                "title": item.title,
                "keywords": item.keywords,
                "product_description": item.product_description,
                "bullet_points": item.bullet_points,
                "price": item.price,
                "image_count": item.image_count,
                "title_rating": item.title_rating,
                "description_rating": item.description_rating,
                "keywords_rating": item.keywords_rating,
                "image_rating": item.image_rating,
                "star_rating": item.star_rating,
                "star_rating_score": star_score,
                "competitor_price": item.competitor_price,
                "competitor_price_rank": competitor_rank,
                "competitor_price_score": competitor_score2,
                "has_ad": has_ad_val,
                "has_aplus": has_aplus_val,
                "has_video": has_video_val,
                "rating_status": item.rating_status or 0,
                "total_score": total_score,
            })

        return {"success": True, "data": result, "total": total}

    total = query.count()
    items = (
        query.order_by(ProductPageInfo.updated_at.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
        .all()
    )

    result = []
    # 查询stores表，建立 inventory_name -> "shop_abbr-site" 映射
    store_rows = db.query(Store.inventory_name, Store.shop_abbr, Store.site).filter(
        Store.tenant_id == current_user.tenant_id,
        Store.inventory_name.isnot(None)
    ).all()
    store_abbr_map = {}
    for r in store_rows:
        if r.inventory_name:
            parts = [p for p in [r.shop_abbr, r.site] if p]
            store_abbr_map[r.inventory_name] = "-".join(parts) if parts else r.inventory_name

    for item in items:
        # 处理 has_aplus：数据库中可能是 None/Boolean/Integer
        has_aplus_val = 0
        if isinstance(item.has_aplus, int):
            has_aplus_val = item.has_aplus
        elif isinstance(item.has_aplus, bool):
            has_aplus_val = 1 if item.has_aplus else 0

        # 处理 has_ad/has_video 为 None 的情况
        has_ad_val = int(item.has_ad) if item.has_ad is not None else None
        has_video_val = bool(item.has_video) if item.has_video is not None else False

        star_score = calc_star_rating_score(item.star_rating)
        competitor_rank, competitor_score = calc_competitor_price_score(item.price, item.competitor_price)
        total_score = calc_total_score(
            item.title_rating,
            item.description_rating,
            item.keywords_rating,
            item.image_rating,
            star_score,
            has_ad_val,
            has_aplus_val,
            has_video_val,
            competitor_score,
        )

        # 店铺名映射：store显示为shop_abbr，保留原始store用于ASIN链接
        display_store = store_abbr_map.get(item.store, item.store) if item.store else item.store

        result.append({
            "id": item.id,
            "tenant_id": item.tenant_id,
            "asin": item.asin,
            "sku": item.sku,
            "store": display_store,
            "store_original": item.store,
            "title": item.title,
            "keywords": item.keywords,
            "product_description": item.product_description,
            "bullet_points": item.bullet_points,
            "price": item.price,
            "image_count": item.image_count,
            "title_rating": item.title_rating,
            "description_rating": item.description_rating,
            "keywords_rating": item.keywords_rating,
            "image_rating": item.image_rating,
            "star_rating": item.star_rating,
            "star_rating_score": star_score,
            "competitor_price": item.competitor_price,
            "competitor_price_rank": competitor_rank,
            "competitor_price_score": competitor_score,
            "has_ad": has_ad_val,
            "has_aplus": has_aplus_val,
            "has_video": has_video_val,
            "rating_status": item.rating_status or 0,
            "total_score": total_score,
        })

    return {"success": True, "data": result, "total": total}


@router.get("/ranking")
async def get_ranking(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """获取评分排行榜（仅已评分记录）"""
    # 查询所有已评分记录
    rated_query = db.query(ProductPageInfo).filter(
        ProductPageInfo.tenant_id == current_user.tenant_id,
        ProductPageInfo.rating_status == 1
    )

    # 非管理员用户按部门过滤数据
    if not is_admin_user(current_user, db):
        user_store_names = get_user_store_names(current_user, db)
        if user_store_names:
            rated_query = rated_query.filter(ProductPageInfo.store.in_(user_store_names))
        else:
            rated_query = rated_query.filter(ProductPageInfo.id == -1)

    rated_items = rated_query.all()

    if not rated_items:
        return {"success": True, "data": {"top10": [], "bottom10": []}}

    # 店铺名映射
    stores = db.query(Store).filter(
        Store.tenant_id == current_user.tenant_id
    ).all()
    store_map = {}
    for s in stores:
        if s.inventory_name:
            parts = [p for p in [s.shop_abbr, s.site] if p]
            store_map[s.inventory_name] = "-".join(parts) if parts else s.inventory_name

    # 计算每个记录的总分
    scored_items = []
    for item in rated_items:
        star_score = calc_star_rating_score(item.star_rating)
        competitor_rank, competitor_score = calc_competitor_price_score(item.price, item.competitor_price)
        has_ad_val = int(item.has_ad) if item.has_ad is not None else 0
        has_aplus_val = item.has_aplus if item.has_aplus is not None else 0
        has_video_val = item.has_video if item.has_video is not None else False

        total_score = calc_total_score(
            item.title_rating,
            item.description_rating,
            item.keywords_rating,
            item.image_rating,
            star_score,
            has_ad_val,
            has_aplus_val,
            has_video_val,
            competitor_score,
        )

        store_display = store_map.get(item.store, item.store)

        scored_items.append({
            "id": item.id,
            "sku": item.sku,
            "asin": item.asin,
            "store": store_display,
            "store_original": item.store,
            "total_score": total_score,
        })

    # 按总分降序排序
    scored_items.sort(key=lambda x: x["total_score"], reverse=True)

    # 取前十
    top10 = scored_items[:10]

    # 取倒数前十
    bottom10 = scored_items[-10:] if len(scored_items) >= 10 else scored_items
    bottom10 = list(reversed(bottom10))

    return {"success": True, "data": {"top10": top10, "bottom10": bottom10}}


@router.get("/{item_id}")
async def get_product_page_info_detail(
    item_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """获取产品页面信息详情"""
    item = db.query(ProductPageInfo).filter(
        ProductPageInfo.id == item_id,
        ProductPageInfo.tenant_id == current_user.tenant_id,
    ).first()

    if not item:
        return {"success": False, "message": "记录不存在"}

    has_aplus_val = 0
    if isinstance(item.has_aplus, int):
        has_aplus_val = item.has_aplus
    elif isinstance(item.has_aplus, bool):
        has_aplus_val = 1 if item.has_aplus else 0

    has_ad_val = int(item.has_ad) if item.has_ad is not None else None
    has_video_val = bool(item.has_video) if item.has_video is not None else False

    star_score = calc_star_rating_score(item.star_rating)
    competitor_rank, competitor_score = calc_competitor_price_score(item.price, item.competitor_price)
    total_score = calc_total_score(
        item.title_rating,
        item.description_rating,
        item.keywords_rating,
        item.image_rating,
        star_score,
        has_ad_val,
        has_aplus_val,
        has_video_val,
        competitor_score,
    )

    # 店铺名映射
    store_rows = db.query(Store.inventory_name, Store.shop_abbr, Store.site).filter(
        Store.tenant_id == current_user.tenant_id,
        Store.inventory_name.isnot(None)
    ).all()
    store_abbr_map = {}
    for r in store_rows:
        if r.inventory_name:
            parts = [p for p in [r.shop_abbr, r.site] if p]
            store_abbr_map[r.inventory_name] = "-".join(parts) if parts else r.inventory_name
    display_store = store_abbr_map.get(item.store, item.store) if item.store else item.store

    return {
        "success": True,
        "data": {
            "id": item.id,
            "tenant_id": item.tenant_id,
            "asin": item.asin,
            "sku": item.sku,
            "store": display_store,
            "store_original": item.store,
            "title": item.title,
            "keywords": item.keywords,
            "product_description": item.product_description,
            "bullet_points": item.bullet_points,
            "price": item.price,
            "image_count": item.image_count,
            "title_rating": item.title_rating,
            "description_rating": item.description_rating,
            "keywords_rating": item.keywords_rating,
            "image_rating": item.image_rating,
            "star_rating": item.star_rating,
            "star_rating_score": star_score,
            "competitor_price": item.competitor_price,
            "competitor_price_rank": competitor_rank,
            "competitor_price_score": competitor_score,
            "has_ad": has_ad_val,
            "has_aplus": has_aplus_val,
            "has_video": has_video_val,
            "total_score": total_score,
            "traffic_keywords": item.traffic_keywords,
            "updated_at": item.updated_at.strftime('%Y-%m-%d %H:%M:%S') if item.updated_at else None,
        },
    }


@router.get("/stores/options")
async def get_store_options(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """获取店铺下拉选项（从 product_page_info 表中提取唯一店铺名，映射为shop_abbr）"""
    stores_query = db.query(ProductPageInfo.store).filter(
        ProductPageInfo.tenant_id == current_user.tenant_id
    )

    # 非管理员用户按部门过滤店铺
    if not is_admin_user(current_user, db):
        user_store_names = get_user_store_names(current_user, db)
        if user_store_names:
            stores_query = stores_query.filter(ProductPageInfo.store.in_(user_store_names))
        else:
            stores_query = stores_query.filter(ProductPageInfo.id == -1)

    stores = stores_query.distinct().all()
    store_names = list(dict.fromkeys([s[0].strip() for s in stores if s[0]]))

    # 查询映射
    store_rows = db.query(Store.inventory_name, Store.shop_abbr, Store.site).filter(
        Store.tenant_id == current_user.tenant_id,
        Store.inventory_name.isnot(None)
    ).all()
    store_abbr_map = {}
    for r in store_rows:
        if r.inventory_name:
            parts = [p for p in [r.shop_abbr, r.site] if p]
            store_abbr_map[r.inventory_name.strip()] = "-".join(parts) if parts else r.inventory_name

    # 按 shop_abbr 去重，只保留在 product_page_info 中有记录的
    # 建立 shop_abbr -> inventory_names 映射
    abbr_to_names = {}
    for r in store_rows:
        if r.inventory_name and r.shop_abbr:
            abbr_to_names.setdefault(r.shop_abbr, []).append(r.inventory_name.strip())

    # 只保留在 store_names 中有匹配的 shop_abbr
    seen_abbrs = set()
    options = []
    for abbr, names in abbr_to_names.items():
        if abbr not in seen_abbrs and any(n in store_names for n in names):
            seen_abbrs.add(abbr)
            options.append({"label": abbr, "value": abbr})
    return {"success": True, "data": options}


class DeleteCompetitorRequest(BaseModel):
    competitor_line: str  # 要删除的竞品行，如 "B0D4TZSF7Y：15.19"


@router.put("/{item_id}/delete-competitor")
async def delete_competitor_price(
    item_id: int,
    request: DeleteCompetitorRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """删除指定竞品价格并重新计算排名和得分"""
    print(f"[DEBUG] 收到删除请求: item_id={item_id}, competitor_line={request.competitor_line}")

    item = db.query(ProductPageInfo).filter(
        ProductPageInfo.id == item_id,
        ProductPageInfo.tenant_id == current_user.tenant_id
    ).first()

    if not item:
        print(f"[DEBUG] 产品不存在: item_id={item_id}")
        return {"success": False, "message": "产品不存在"}

    print(f"[DEBUG] 当前竞品价格数据: {item.competitor_price}")

    # 解析现有竞品价格
    competitor_lines = []
    if item.competitor_price:
        for line in item.competitor_price.strip().splitlines():
            line = line.strip()
            if line:
                competitor_lines.append(line)

    print(f"[DEBUG] 解析后的竞品列表: {competitor_lines}")

    # 清理要删除的竞品行（去除可能的换行符）
    target_line = request.competitor_line.strip().replace('\r', '')
    print(f"[DEBUG] 目标删除行: {target_line}")

    # 删除指定的竞品行
    new_lines = [l for l in competitor_lines if l != target_line]
    print(f"[DEBUG] 删除后的竞品列表: {new_lines}")

    # 更新数据库
    item.competitor_price = "\n".join(new_lines) if new_lines else None

    # 重新计算排名和得分
    competitor_rank, competitor_score = calc_competitor_price_score(item.price, item.competitor_price)

    # 处理 has_aplus
    has_aplus_val = 0
    if isinstance(item.has_aplus, int):
        has_aplus_val = item.has_aplus
    elif isinstance(item.has_aplus, bool):
        has_aplus_val = 1 if item.has_aplus else 0

    star_score = calc_star_rating_score(item.star_rating)
    total_score = calc_total_score(
        item.title_rating,
        item.description_rating,
        item.keywords_rating,
        item.image_rating,
        star_score,
        item.has_ad,
        has_aplus_val,
        item.has_video,
        competitor_score,
    )

    db.commit()
    print(f"[DEBUG] 删除成功，新排名: {competitor_rank}, 新得分: {competitor_score}")

    return {
        "success": True,
        "data": {
            "competitor_price": item.competitor_price,
            "competitor_price_rank": competitor_rank,
            "competitor_price_score": competitor_score,
            "total_score": total_score,
        }
    }


class SubmitRatingRequest(BaseModel):
    ids: List[int]  # 选中的产品ID列表


@router.post("/submit-rating")
async def submit_rating(
    request: SubmitRatingRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """提交评分到飞书webhook"""
    if not request.ids:
        return {"success": False, "message": "未选择任何数据"}

    # 查询选中的产品数据
    items = db.query(ProductPageInfo).filter(
        ProductPageInfo.id.in_(request.ids),
        ProductPageInfo.tenant_id == current_user.tenant_id
    ).all()

    if not items:
        return {"success": False, "message": "未找到选中数据"}

    # 构造飞书请求格式
    feishu_data_list = []
    for item in items:
        feishu_data_list.append({
            "sku": item.sku or "",
            "asin": item.asin or "",
            "店铺": item.store or "",
            "标题": item.title or "",
            "五点描述": item.bullet_points or "",
            "关键词": item.keywords or "",
            "图片数": str(item.image_count) if item.image_count else "0",
            "产品描述": item.product_description or "",
            "价格": item.price or "",
            "数据库ID": str(item.id),
        })

    # 发送到飞书webhook
    feishu_webhook_url = "https://pcn4p6l5do51.feishu.cn/base/automation/webhook/event/KgcgaTUPhwqeE9hB09CciwyLnrb"

    try:
        async with httpx.AsyncClient() as client:
            # 每条数据单独发送
            for data in feishu_data_list:
                response = await client.post(
                    feishu_webhook_url,
                    json=data,
                    headers={"Content-Type": "application/json"},
                    timeout=10.0
                )
                print(f"[DEBUG] 飞书响应: {response.status_code}, 数据ID: {data['数据库ID']}")

        # 提交成功后，将选中记录的评分状态改为未评分(0)
        db.query(ProductPageInfo).filter(
            ProductPageInfo.id.in_(request.ids),
            ProductPageInfo.tenant_id == current_user.tenant_id
        ).update({"rating_status": 0}, synchronize_session="fetch")
        db.commit()

        return {
            "success": True,
            "message": f"成功提交 {len(feishu_data_list)} 条数据",
            "data": {"count": len(feishu_data_list)}
        }
    except Exception as e:
        print(f"[DEBUG] 飞书请求失败: {str(e)}")
        return {"success": False, "message": f"发送到飞书失败: {str(e)}"}


@router.post("/delete-records")
async def delete_records(
    request: SubmitRatingRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """删除选中的记录"""
    if not request.ids:
        return {"success": False, "message": "未选择任何数据"}

    deleted = db.query(ProductPageInfo).filter(
        ProductPageInfo.id.in_(request.ids),
        ProductPageInfo.tenant_id == current_user.tenant_id
    ).delete(synchronize_session="fetch")
    db.commit()

    return {
        "success": True,
        "message": f"成功删除 {deleted} 条数据",
        "data": {"count": deleted}
    }


@router.post("/import-excel")
async def import_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """导入Excel表格，读取SKU和店铺列新增到数据库"""
    if not file.filename.endswith(('.xlsx', '.xls')):
        return {"success": False, "message": "只支持Excel表格文件（.xlsx/.xls）"}

    try:
        contents = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
        ws = wb.active

        # 读取表头，找到SKU和店铺列的索引
        headers = [str(cell.value).strip() if cell.value else '' for cell in ws[1]]
        sku_col = None
        store_col = None

        for i, h in enumerate(headers):
            if h.upper() == 'SKU':
                sku_col = i
            elif h == '店铺':
                store_col = i

        if sku_col is None:
            return {"success": False, "message": "未找到SKU列，请确保表头包含'SKU'"}
        if store_col is None:
            return {"success": False, "message": "未找到店铺列，请确保表头包含'店铺'"}

        # 查询当前租户下已存在的(sku, store)组合，用于去重
        existing = db.query(ProductPageInfo.sku, ProductPageInfo.store).filter(
            ProductPageInfo.tenant_id == current_user.tenant_id
        ).all()
        existing_set = {(r[0], r[1]) for r in existing}

        # 批量构建数据
        records = []
        skipped_count = 0
        duplicate_count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) <= max(sku_col, store_col):
                skipped_count += 1
                continue

            sku_val = str(row[sku_col]).strip() if row[sku_col] else ''
            store_val = str(row[store_col]).strip() if row[store_col] else ''

            if not sku_val:
                skipped_count += 1
                continue

            # 去重：sku和店铺都相同时跳过
            if (sku_val, store_val) in existing_set:
                duplicate_count += 1
                continue

            records.append({
                "tenant_id": current_user.tenant_id,
                "sku": sku_val,
                "store": store_val,
                "rating_status": 0,
            })

        wb.close()

        if not records:
            msg = "未读取到有效数据"
            if duplicate_count > 0:
                msg += f"，已跳过 {duplicate_count} 条重复数据"
            return {"success": False, "message": msg}

        # 批量插入，全部插入后再统一提交，支持事务回滚
        added_count = 0
        new_ids = []
        for record in records:
            item = ProductPageInfo(**record)
            db.add(item)
            db.flush()  # flush获取id但不提交
            new_ids.append(item.id)
            added_count += 1

        db.commit()

        msg = f"成功导入 {added_count} 条数据"
        if duplicate_count > 0:
            msg += f"，跳过 {duplicate_count} 条重复数据"
        if skipped_count > 0:
            msg += f"，无效行 {skipped_count} 条"

        return {
            "success": True,
            "message": msg,
            "data": {"added": added_count, "skipped": skipped_count, "duplicate": duplicate_count, "import_ids": new_ids}
        }
    except Exception as e:
        db.rollback()
        print(f"[DEBUG] 导入Excel失败: {str(e)}")
        import traceback
        traceback.print_exc()
        return {"success": False, "message": f"导入失败: {str(e)}"}


@router.post("/cancel-import")
async def cancel_import(
    request: Request,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """取消最近一次导入，删除本次导入的所有记录"""
    body = await request.json()
    import_ids = body.get("import_ids", [])
    if not import_ids:
        return {"success": False, "message": "没有可取消的导入记录"}

    try:
        deleted = db.query(ProductPageInfo).filter(
            ProductPageInfo.id.in_(import_ids),
            ProductPageInfo.tenant_id == current_user.tenant_id
        ).delete(synchronize_session=False)
        db.commit()
        return {"success": True, "message": f"已取消导入，删除 {deleted} 条记录"}
    except Exception as e:
        db.rollback()
        return {"success": False, "message": f"取消导入失败: {str(e)}"}


@router.post("/submit-edit")
async def submit_edit(
    request: Request,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """提交编辑基础信息 - 发送飞书webhook并可选更新评分状态"""
    body = await request.json()
    item_id = body.get("item_id")
    store = body.get("store", "")
    field_type = body.get("field_type", "")
    sku = body.get("sku", "")
    content = body.get("content", "")
    reset_rating = body.get("reset_rating", False)

    if not item_id:
        return {"success": False, "message": "缺少item_id参数"}

    # 验证记录存在且属于当前租户
    item = db.query(ProductPageInfo).filter(
        ProductPageInfo.id == item_id,
        ProductPageInfo.tenant_id == current_user.tenant_id,
    ).first()

    if not item:
        return {"success": False, "message": "记录不存在"}

    # 发送飞书webhook
    try:
        async with httpx.AsyncClient() as client:
            await client.post(
                "https://pcn4p6l5do51.feishu.cn/base/automation/webhook/event/AnnPaHjm1wTDvohfMnDc4A2Wnlg",
                json={
                    "店铺": store,
                    "选择填写类型": field_type,
                    "SKU": sku,
                    "填写文本": content,
                },
                timeout=10,
            )
    except Exception as e:
        logger.error(f"飞书webhook发送失败: {e}")
        return {"success": False, "message": f"飞书通知发送失败: {str(e)}"}

    # 如果需要重新评分，更新评分状态
    if reset_rating:
        item.rating_status = 0
        db.commit()

    return {"success": True, "message": "提交成功"}


@router.put("/{item_id}/rating-status")
async def update_rating_status(
    item_id: int,
    request: Request,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """更新评分状态"""
    body = await request.json()
    rating_status = body.get("rating_status")
    if rating_status is None:
        return {"success": False, "message": "缺少rating_status参数"}

    item = db.query(ProductPageInfo).filter(
        ProductPageInfo.id == item_id,
        ProductPageInfo.tenant_id == current_user.tenant_id,
    ).first()

    if not item:
        return {"success": False, "message": "记录不存在"}

    item.rating_status = rating_status
    db.commit()
    return {"success": True, "message": "评分状态已更新"}