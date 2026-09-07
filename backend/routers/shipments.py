from fastapi import APIRouter, HTTPException, Depends, Query, UploadFile, File
from pydantic import BaseModel
from typing import List, Optional
from sqlalchemy.orm import Session
from sqlalchemy import text
from datetime import datetime
import io
import openpyxl

from database.database import get_db
from dependencies import get_current_user, PermissionChecker
from models.user import User
from services.operation_log import log_order_create, log_order_confirm, log_order_delete, log_order_update

router = APIRouter(prefix="/api/shipments", tags=["shipments"])


class ShipmentItemCreate(BaseModel):
    product_id: Optional[int] = None
    product_code: Optional[str] = None
    product_name: Optional[str] = None
    stock_quantity: Optional[int] = 0
    red_list: Optional[str] = None
    sea_freight: Optional[str] = None
    notes: Optional[str] = None
    sku: Optional[str] = None


class ShipmentCreate(BaseModel):
    order_number: str
    store_group_id: Optional[int] = None
    store_group_name: Optional[str] = None
    notes: Optional[str] = None
    items: List[ShipmentItemCreate]


class ShipmentUpdate(BaseModel):
    order_number: Optional[str] = None
    store_group_id: Optional[int] = None
    store_group_name: Optional[str] = None
    notes: Optional[str] = None
    items: Optional[List[ShipmentItemCreate]] = None


class ConvertToOutboundRequest(BaseModel):
    ids: List[int]
    notes: Optional[str] = None


@router.get("/")
async def get_shipments(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    status: Optional[str] = None,
    search: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    store_group_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    try:
        where_conditions = ["so.tenant_id = :tenant_id", "so.deleted_at IS NULL"]
        params = {"tenant_id": current_user.tenant_id}

        if status:
            where_conditions.append("so.status = :status")
            params["status"] = status
        if search:
            where_conditions.append("(so.order_number LIKE :search OR so.store_group_name LIKE :search OR so.notes LIKE :search)")
            params["search"] = f"%{search}%"
        if start_date:
            where_conditions.append("DATE(so.created_at) >= :start_date")
            params["start_date"] = start_date
        if end_date:
            where_conditions.append("DATE(so.created_at) <= :end_date")
            params["end_date"] = end_date
        if store_group_id:
            where_conditions.append("so.store_group_id = :store_group_id")
            params["store_group_id"] = store_group_id

        where_clause = " AND ".join(where_conditions)

        total = db.execute(text(f"SELECT COUNT(*) FROM shipment_orders so WHERE {where_clause}"), params).scalar() or 0

        offset = (page - 1) * page_size
        params["limit"] = page_size
        params["offset"] = offset

        rows = db.execute(text(f"""
            SELECT so.id, so.order_number, so.store_group_id, so.store_group_name,
                   so.total_quantity, so.status, so.notes,
                   so.created_by, so.confirmed_by, so.confirmed_at, so.created_at,
                   so.outbound_order_id, ob.order_number as outbound_order_number
            FROM shipment_orders so
            LEFT JOIN outbound_orders ob ON ob.id = so.outbound_order_id
            WHERE {where_clause}
            ORDER BY so.created_at DESC
            LIMIT :limit OFFSET :offset
        """), params).fetchall()

        user_ids = set()
        for row in rows:
            if row[7]:
                user_ids.add(row[7])
            if row[8]:
                user_ids.add(row[8])

        user_map = {}
        if user_ids:
            user_rows = db.execute(text(
                "SELECT id, nickname, username FROM users WHERE id IN :ids"
            ), {"ids": tuple(user_ids)}).fetchall()
            for u in user_rows:
                user_map[u[0]] = u[1] or u[2]

        result = []
        for row in rows:
            result.append({
                "id": row[0],
                "order_number": row[1],
                "store_group_id": row[2],
                "store_group_name": row[3] or "",
                "total_quantity": int(row[4]) if row[4] else 0,
                "status": row[5],
                "notes": row[6],
                "created_by": row[7],
                "creator_name": user_map.get(row[7], "") if row[7] else "",
                "confirmed_by": row[8],
                "confirmer_name": user_map.get(row[8], "") if row[8] else "",
                "confirmed_at": row[9].strftime("%Y-%m-%d %H:%M:%S") if row[9] else None,
                "created_at": row[10].strftime("%Y-%m-%d %H:%M:%S") if row[10] else "",
                "outbound_order_id": row[11],
                "outbound_order_number": row[12] or None,
            })

        return {"success": True, "data": result, "total": total, "page": page, "page_size": page_size}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"获取发货单失败: {str(e)}")


@router.post("/")
async def create_shipment(
    data: ShipmentCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(PermissionChecker("shipment:create"))
):
    try:
        order_number = f"FH{datetime.now().strftime('%Y%m%d%H%M%S')}"
        total_qty = sum((item.stock_quantity or 0) for item in data.items)
        creator_name = current_user.nickname or current_user.username

        db.execute(text("""
            INSERT INTO shipment_orders (tenant_id, order_number, store_group_id, store_group_name,
                total_quantity, status, notes,
                created_by, creator_name, created_at, updated_at)
            VALUES (:tid, :num, :sg_id, :sg_name, :tq, 'draft', :notes, :uid, :cname, NOW(), NOW())
        """), {
            "tid": current_user.tenant_id, "num": order_number,
            "sg_id": data.store_group_id, "sg_name": data.store_group_name,
            "tq": total_qty, "notes": data.notes,
            "uid": current_user.id, "cname": creator_name,
        })
        order_id = db.execute(text("SELECT LAST_INSERT_ID()")).scalar()

        # 获取当前用户的店铺ID列表，用于自动匹配SKU
        user_store_rows = db.execute(text("""
            SELECT us.store_id FROM user_stores us
            JOIN stores s ON us.store_id = s.id
            WHERE us.user_id = :uid AND us.tenant_id = :tid AND s.deleted_at IS NULL
        """), {"uid": current_user.id, "tid": current_user.tenant_id}).fetchall()
        user_store_ids = [r[0] for r in user_store_rows]

        for item in data.items:
            pid = item.product_id
            if pid is None and item.product_code:
                product = db.execute(text(
                    "SELECT id FROM products WHERE product_code = :code AND tenant_id = :tid AND deleted_at IS NULL"
                ), {"code": item.product_code, "tid": current_user.tenant_id}).fetchone()
                if product:
                    pid = product[0]

            # 自动获取SKU：如果前端没传SKU，根据当前用户的店铺匹配
            sku = item.sku
            if not sku and pid and user_store_ids:
                store_placeholders = ",".join([f":us_{i}" for i in range(len(user_store_ids))])
                sku_params = {"pid": pid, "tid": current_user.tenant_id}
                for i, sid in enumerate(user_store_ids):
                    sku_params[f"us_{i}"] = sid
                sku_row = db.execute(text(f"""
                    SELECT pp.sku FROM platform_products pp
                    WHERE pp.product_id = :pid AND pp.deleted_at IS NULL AND pp.tenant_id = :tid
                    AND (
                        pp.store_id IN ({store_placeholders})
                        OR EXISTS (
                            SELECT 1 FROM ({' UNION ALL '.join([f'SELECT :us_{i} as sid' for i in range(len(user_store_ids))])}) as u
                            WHERE JSON_CONTAINS(pp.store_id, CAST(u.sid AS JSON))
                        )
                    )
                    LIMIT 1
                """), sku_params).fetchone()
                if sku_row:
                    sku = sku_row[0]

            # 如果发货单有store_group_id，也可以通过store_group匹配SKU
            if not sku and pid and data.store_group_id:
                sku_row = db.execute(text("""
                    SELECT pp.sku FROM platform_products pp
                    WHERE pp.product_id = :pid AND pp.deleted_at IS NULL AND pp.tenant_id = :tid
                    AND EXISTS (
                        SELECT 1 FROM stores s
                        WHERE s.group_id = :sg_id AND s.deleted_at IS NULL
                        AND JSON_CONTAINS(pp.store_id, CAST(s.id AS JSON))
                    )
                    LIMIT 1
                """), {"pid": pid, "tid": current_user.tenant_id, "sg_id": data.store_group_id}).fetchone()
                if sku_row:
                    sku = sku_row[0]

            db.execute(text("""
                INSERT INTO shipment_order_items (tenant_id, shipment_order_id, product_id,
                    product_code, product_name, stock_quantity, sku, created_at, updated_at)
                VALUES (:tid, :oid, :pid, :pcode, :pname, :sqty, :sku, NOW(), NOW())
            """), {
                "tid": current_user.tenant_id, "oid": order_id, "pid": pid,
                "pcode": item.product_code, "pname": item.product_name,
                "sqty": item.stock_quantity or 0, "sku": sku or "",
            })

        db.commit()

        log_order_create(db, current_user.tenant_id, current_user.id, creator_name,
                         "shipment", order_id, order_number,
                         {"store_group_id": data.store_group_id,
                          "store_group_name": data.store_group_name or "",
                          "items_count": len(data.items)})
        db.commit()

        return {"success": True, "id": order_id, "order_number": order_number}
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"创建发货单失败: {str(e)}")


@router.get("/kpi-count")
async def get_kpi_count(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """获取各状态发货单数量（KPI卡片）"""
    try:
        rows = db.execute(text("""
            SELECT status, COUNT(*) as cnt
            FROM shipment_orders
            WHERE tenant_id = :tid AND deleted_at IS NULL
            GROUP BY status
        """), {"tid": current_user.tenant_id}).fetchall()

        result = {
            "pending_shipments": 0,      # 草稿
            "confirmed_shipments": 0,     # 已确认
            "cancelled_shipments": 0,     # 已取消
        }
        for row in rows:
            status = row[0]
            cnt = row[1]
            if status == 'draft':
                result["pending_shipments"] = cnt
            elif status == 'confirmed':
                result["confirmed_shipments"] = cnt
            elif status == 'cancelled':
                result["cancelled_shipments"] = cnt

        return {"success": True, **result}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"获取发货单KPI数量失败: {str(e)}")


@router.get("/{order_id}")
async def get_shipment_detail(
    order_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    try:
        order = db.execute(text("""
            SELECT so.id, so.order_number, so.store_group_id, so.store_group_name,
                   so.total_quantity, so.status, so.notes,
                   so.created_by, so.confirmed_by, so.confirmed_at, so.created_at
            FROM shipment_orders so
            WHERE so.id = :id AND so.tenant_id = :tid AND so.deleted_at IS NULL
        """), {"id": order_id, "tid": current_user.tenant_id}).fetchone()

        if not order:
            raise HTTPException(status_code=404, detail="发货单不存在")

        items = db.execute(text("""
            SELECT soi.id, soi.product_id, soi.product_code, soi.product_name,
                   soi.stock_quantity,
                   soi.red_list, soi.sea_freight, soi.notes,
                   CASE 
                       WHEN soi.sku IS NOT NULL AND soi.sku != '' THEN soi.sku
                       ELSE (
                           SELECT pp.sku FROM platform_products pp 
                           JOIN stores s ON s.group_id = :sg_id AND s.deleted_at IS NULL
                           WHERE pp.product_id = soi.product_id 
                           AND pp.deleted_at IS NULL
                           AND pp.tenant_id = :tid
                           AND (pp.store_id = s.id OR JSON_CONTAINS(pp.store_id, CAST(s.id AS JSON)))
                           LIMIT 1
                       )
                   END as platform_sku
            FROM shipment_order_items soi
            WHERE soi.shipment_order_id = :oid AND soi.deleted_at IS NULL
        """), {"oid": order_id, "tid": current_user.tenant_id, "sg_id": order[2]}).fetchall()

        item_list = []
        for it in items:
            item_list.append({
                "id": it[0],
                "product_id": it[1],
                "product_code": it[2] or "",
                "product_name": it[3] or "",
                "stock_quantity": int(it[4]) if it[4] else 0,
                "red_list": it[5] or "",
                "sea_freight": it[6] or "",
                "notes": it[7] or "",
                "sku": it[8] or "",
            })

        user_ids = set()
        if order[7]:
            user_ids.add(order[7])
        if order[8]:
            user_ids.add(order[8])
        user_map = {}
        if user_ids:
            user_rows = db.execute(text(
                "SELECT id, nickname, username FROM users WHERE id IN :ids"
            ), {"ids": tuple(user_ids)}).fetchall()
            for u in user_rows:
                user_map[u[0]] = u[1] or u[2]

        return {
            "id": order[0],
            "order_number": order[1],
            "store_group_id": order[2],
            "store_group_name": order[3] or "",
            "total_quantity": int(order[4]) if order[4] else 0,
            "status": order[5],
            "notes": order[6] or "",
            "created_by": order[7],
            "creator_name": user_map.get(order[7], "") if order[7] else "",
            "confirmed_by": order[8],
            "confirmer_name": user_map.get(order[8], "") if order[8] else "",
            "confirmed_at": order[9].strftime("%Y-%m-%d %H:%M:%S") if order[9] else None,
            "created_at": order[10].strftime("%Y-%m-%d %H:%M:%S") if order[10] else "",
            "items": item_list,
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"获取发货单详情失败: {str(e)}")


@router.put("/{order_id}")
async def update_shipment(
    order_id: int,
    data: ShipmentUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(PermissionChecker("shipment:edit"))
):
    try:
        order = db.execute(text(
            "SELECT id, status, order_number, store_group_id, store_group_name FROM shipment_orders WHERE id = :id AND tenant_id = :tid AND deleted_at IS NULL"
        ), {"id": order_id, "tid": current_user.tenant_id}).fetchone()
        if not order:
            raise HTTPException(status_code=404, detail="发货单不存在")
        if order[1] != "draft":
            raise HTTPException(status_code=400, detail="只有草稿状态的发货单可以编辑")

        before_data = {
            "order_number": order[2],
            "store_group_id": order[3],
            "store_group_name": order[4],
        }

        set_parts = ["updated_at = NOW()"]
        params = {"id": order_id}
        order_number = order[2]

        if data.order_number is not None:
            set_parts.append("order_number = :num")
            params["num"] = data.order_number
            order_number = data.order_number
        if data.store_group_id is not None:
            set_parts.append("store_group_id = :sg_id")
            params["sg_id"] = data.store_group_id
        if data.store_group_name is not None:
            set_parts.append("store_group_name = :sg_name")
            params["sg_name"] = data.store_group_name
        if data.notes is not None:
            set_parts.append("notes = :notes")
            params["notes"] = data.notes

        # 如果有items，先软删除旧items再插入新items
        if data.items:
            db.execute(text(
                "UPDATE shipment_order_items SET deleted_at = NOW() WHERE shipment_order_id = :oid AND deleted_at IS NULL"
            ), {"oid": order_id})

            total_qty = sum((item.stock_quantity or 0) for item in data.items)
            set_parts.append("total_quantity = :tq")
            params["tq"] = total_qty

            for item in data.items:
                pid = item.product_id
                if pid is None and item.product_code:
                    product = db.execute(text(
                        "SELECT id FROM products WHERE product_code = :code AND tenant_id = :tid AND deleted_at IS NULL"
                    ), {"code": item.product_code, "tid": current_user.tenant_id}).fetchone()
                    if product:
                        pid = product[0]
                db.execute(text("""
                    INSERT INTO shipment_order_items (tenant_id, shipment_order_id, product_id,
                        product_code, product_name, stock_quantity,
                        red_list, sea_freight, notes, sku, created_at, updated_at)
                    VALUES (:tid, :oid, :pid, :pcode, :pname, :sqty,
                        :red_list, :sea_freight, :notes, :sku, NOW(), NOW())
                """), {
                    "tid": current_user.tenant_id, "oid": order_id, "pid": pid,
                    "pcode": item.product_code, "pname": item.product_name,
                    "sqty": item.stock_quantity or 0,
                    "red_list": item.red_list,
                    "sea_freight": item.sea_freight,
                    "notes": item.notes,
                    "sku": item.sku or "",
                })

        db.execute(text(f"UPDATE shipment_orders SET {', '.join(set_parts)} WHERE id = :id"), params)
        db.commit()

        after_data = {
            "order_number": order_number,
            "store_group_id": data.store_group_id if data.store_group_id is not None else order[3],
            "store_group_name": data.store_group_name if data.store_group_name is not None else order[4],
        }

        log_order_update(db, current_user.tenant_id, current_user.id,
                         current_user.nickname or current_user.username,
                         "shipment", order_id, order_number,
                         before_data, after_data)
        db.commit()

        return {"success": True}
    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"更新发货单失败: {str(e)}")


@router.put("/{order_id}/confirm")
async def confirm_shipment(
    order_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(PermissionChecker("shipment:confirm"))
):
    try:
        order = db.execute(text(
            "SELECT id, order_number, status FROM shipment_orders WHERE id = :id AND tenant_id = :tid AND deleted_at IS NULL"
        ), {"id": order_id, "tid": current_user.tenant_id}).fetchone()
        if not order:
            raise HTTPException(status_code=404, detail="发货单不存在")
        if order[2] != "draft":
            raise HTTPException(status_code=400, detail=f"当前状态'{order[2]}'不允许确认")

        confirmer_name = current_user.nickname or current_user.username
        db.execute(text("""
            UPDATE shipment_orders SET
                status = 'confirmed',
                confirmed_by = :uid,
                confirmer_name = :cname,
                confirmed_at = NOW(),
                updated_at = NOW()
            WHERE id = :id
        """), {"uid": current_user.id, "cname": confirmer_name, "id": order_id})
        db.commit()

        log_order_confirm(db, current_user.tenant_id, current_user.id, confirmer_name,
                          "shipment", order_id, order[1],
                          {"status": "draft"},
                          {"status": "confirmed"})
        db.commit()

        return {"success": True}
    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"确认发货单失败: {str(e)}")


@router.delete("/{order_id}")
async def delete_shipment(
    order_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(PermissionChecker("shipment:delete"))
):
    try:
        order = db.execute(text(
            "SELECT id, order_number, status FROM shipment_orders WHERE id = :id AND tenant_id = :tid AND deleted_at IS NULL"
        ), {"id": order_id, "tid": current_user.tenant_id}).fetchone()
        if not order:
            raise HTTPException(status_code=404, detail="发货单不存在")

        order_status = order[2]

        before_data = {"order_number": order[1], "status": order_status}

        db.execute(text(
            "UPDATE shipment_order_items SET deleted_at = NOW() WHERE shipment_order_id = :oid AND deleted_at IS NULL"
        ), {"oid": order_id})
        db.execute(text(
            "UPDATE shipment_orders SET deleted_at = NOW() WHERE id = :id"
        ), {"id": order_id})
        db.commit()

        log_order_delete(db, current_user.tenant_id, current_user.id,
                         current_user.nickname or current_user.username,
                         "shipment", order_id, order[1], before_data)
        db.commit()

        return {"success": True}
    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"删除发货单失败: {str(e)}")


@router.post("/batch-confirm")
async def batch_confirm_shipments(
    data: dict,
    db: Session = Depends(get_db),
    current_user: User = Depends(PermissionChecker("shipment:confirm"))
):
    """批量确认发货单"""
    try:
        ids = data.get("ids", [])
        if not ids:
            raise HTTPException(status_code=400, detail="请选择要确认的发货单")

        confirmer_name = current_user.nickname or current_user.username
        placeholders = ",".join([f":id_{i}" for i in range(len(ids))])
        params = {"tid": current_user.tenant_id, "cname": confirmer_name}
        for i, oid in enumerate(ids):
            params[f"id_{i}"] = oid

        result = db.execute(text(f"""
            UPDATE shipment_orders
            SET status = 'confirmed', confirmed_by = :uid, confirmer_name = :cname, confirmed_at = NOW(), updated_at = NOW()
            WHERE id IN ({placeholders})
            AND tenant_id = :tid AND status = 'draft' AND deleted_at IS NULL
        """), {**params, "uid": current_user.id})

        db.commit()
        confirmed_count = result.rowcount

        for oid in ids:
            log_order_confirm(db, current_user.tenant_id, current_user.id, confirmer_name,
                              "shipment", oid, "", {})
        db.commit()

        return {"success": True, "confirmed_count": confirmed_count}
    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"批量确认发货单失败: {str(e)}")


@router.post("/batch-delete")
async def batch_delete_shipments(
    data: dict,
    db: Session = Depends(get_db),
    current_user: User = Depends(PermissionChecker("shipment:delete"))
):
    """批量删除发货单"""
    try:
        ids = data.get("ids", [])
        if not ids:
            raise HTTPException(status_code=400, detail="请选择要删除的发货单")

        placeholders = ",".join([f":id_{i}" for i in range(len(ids))])
        params = {"tid": current_user.tenant_id}
        for i, oid in enumerate(ids):
            params[f"id_{i}"] = oid

        result = db.execute(text(f"""
            UPDATE shipment_orders SET deleted_at = NOW()
            WHERE id IN ({placeholders})
            AND tenant_id = :tid AND deleted_at IS NULL
        """), params)

        db.commit()
        deleted_count = result.rowcount

        return {"success": True, "deleted_count": deleted_count}
    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"批量删除发货单失败: {str(e)}")


@router.post("/batch-convert-outbound")
async def batch_convert_to_outbound(
    data: ConvertToOutboundRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(PermissionChecker("shipment:edit"))
):
    """批量将发货单转为出库单，同店铺分组的发货单合并为同一张出库单"""
    try:
        if not data.ids:
            raise HTTPException(status_code=400, detail="请至少选择一条发货单")

        # 1. 查询选中的发货单（状态必须是 confirmed 且未关联出库单）
        ids_placeholders = ', '.join(f':id{i}' for i in range(len(data.ids)))
        order_params = {f'id{i}': data.ids[i] for i in range(len(data.ids))}
        order_params["tenant_id"] = current_user.tenant_id

        orders = db.execute(text(f"""
            SELECT so.id, so.order_number, so.store_group_id, so.store_group_name
            FROM shipment_orders so
            WHERE so.id IN ({ids_placeholders}) AND so.tenant_id = :tenant_id
              AND so.deleted_at IS NULL AND so.status = 'confirmed' AND so.outbound_order_id IS NULL
            ORDER BY so.created_at ASC
        """), order_params).fetchall()

        if not orders:
            raise HTTPException(status_code=400, detail="未找到可转换的发货单（需为已确认且未关联出库单）")

        # 2. 按店铺分组分组
        group_groups = {}  # store_group_id -> [(id, order_number, group_name), ...]
        for o in orders:
            group_id = o[2] or 0
            group_name = o[3] or "未分组"
            if group_id not in group_groups:
                group_groups[group_id] = []
            group_groups[group_id].append((o[0], o[1], group_name))

        # 3. 自动选择最新创建的活跃仓库
        warehouse = None
        latest_wh = db.execute(text("""
            SELECT name FROM warehouses
            WHERE tenant_id = :tid AND deleted_at IS NULL AND status = 'active'
            ORDER BY created_at DESC LIMIT 1
        """), {"tid": current_user.tenant_id}).fetchone()
        if latest_wh:
            warehouse = latest_wh[0]

        now = datetime.now()
        created_ob_numbers = []
        ob_seq = 0

        for store_group_id, group_orders in group_groups.items():
            group_order_ids = [o[0] for o in group_orders]
            group_name = group_orders[0][2] if group_orders else "未分组"

            # 查询该组发货单的明细
            oi_placeholders = ', '.join(f':oid{i}' for i in range(len(group_order_ids)))
            item_params = {f'oid{i}': group_order_ids[i] for i in range(len(group_order_ids))}

            items = db.execute(text(f"""
                SELECT soi.product_id, soi.product_code, soi.product_name, soi.stock_quantity
                FROM shipment_order_items soi
                WHERE soi.shipment_order_id IN ({oi_placeholders}) AND soi.deleted_at IS NULL
            """), item_params).fetchall()

            if not items:
                continue

            # 按产品ID汇总数量
            product_agg = {}
            for item in items:
                pid = item[0]
                qty = int(item[3])
                if pid in product_agg:
                    product_agg[pid]["quantity"] += qty
                else:
                    product_agg[pid] = {
                        "product_id": pid,
                        "product_code": item[1] or "",
                        "product_name": item[2] or "",
                        "quantity": qty,
                    }

            # 查询产品采购价
            all_pids = list(product_agg.keys())
            price_map = {}
            if all_pids:
                pid_placeholders = ', '.join(f':pid{i}' for i in range(len(all_pids)))
                pid_params = {f'pid{i}': all_pids[i] for i in range(len(all_pids))}
                price_rows = db.execute(text(f"""
                    SELECT id, purchase_price FROM products WHERE id IN ({pid_placeholders})
                """), pid_params).fetchall()
                for pr in price_rows:
                    price_map[pr[0]] = float(pr[1]) if pr[1] else 0.0

            total_qty = sum(agg["quantity"] for agg in product_agg.values())
            total_amt = sum(agg["quantity"] * price_map.get(pid, 0.0) for pid, agg in product_agg.items())

            # 创建出库单
            ob_seq += 1
            ob_number = f"OB{now.strftime('%Y%m%d%H%M%S')}{ob_seq:02d}"
            db.execute(text("""
                INSERT INTO outbound_orders (tenant_id, order_number, outbound_type,
                    warehouse, outbound_date, total_quantity, total_amount, status, store_group_id, notes, created_by, created_at, updated_at)
                VALUES (:tenant_id, :order_number, 'shipment_fba',
                    :warehouse, :outbound_date, :total_quantity, :total_amount, 'draft', :store_group_id, :notes, :created_by, :created_at, :updated_at)
            """), {
                "tenant_id": current_user.tenant_id,
                "order_number": ob_number,
                "warehouse": warehouse,
                "outbound_date": now,
                "total_quantity": total_qty,
                "total_amount": total_amt,
                "store_group_id": store_group_id if store_group_id != 0 else None,
                "notes": data.notes or "",
                "created_by": current_user.id,
                "created_at": now,
                "updated_at": now,
            })
            ob_id = db.execute(text("SELECT LAST_INSERT_ID()")).scalar()

            # 创建出库单明细
            for pid, agg in product_agg.items():
                unit_price = price_map.get(pid, 0.0)
                total_price = agg["quantity"] * unit_price
                db.execute(text("""
                    INSERT INTO outbound_order_items (outbound_order_id, product_id, quantity, unit_price, total_price, notes, created_at, updated_at)
                    VALUES (:oid, :pid, :qty, :up, :tp, :notes, :created_at, :updated_at)
                """), {
                    "oid": ob_id, "pid": pid, "qty": agg["quantity"],
                    "up": unit_price, "tp": total_price, "notes": "",
                    "created_at": now, "updated_at": now,
                })

            # 更新发货单关联出库单ID
            for order_id, _, _ in group_orders:
                db.execute(text("""
                    UPDATE shipment_orders
                    SET outbound_order_id = :ob_id, updated_at = :updated_at
                    WHERE id = :id
                """), {
                    "ob_id": ob_id,
                    "updated_at": now,
                    "id": order_id,
                })

            created_ob_numbers.append(ob_number)

        db.commit()

        return {
            "success": True,
            "message": f"已将 {len(orders)} 条发货单转为 {len(created_ob_numbers)} 张出库单",
            "data": {
                "outbound_order_numbers": created_ob_numbers,
                "converted_count": len(orders),
                "ob_count": len(created_ob_numbers),
            }
        }
    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"批量转出库单失败: {str(e)}")


@router.get("/{order_id}/export")
async def export_shipment_detail(
    order_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """导出发货单明细为Excel"""
    try:
        order = db.execute(text("""
            SELECT so.id, so.order_number, so.store_group_name
            FROM shipment_orders so
            WHERE so.id = :id AND so.tenant_id = :tid AND so.deleted_at IS NULL
        """), {"id": order_id, "tid": current_user.tenant_id}).fetchone()

        if not order:
            raise HTTPException(status_code=404, detail="发货单不存在")

        items = db.execute(text("""
            SELECT soi.id, soi.product_code, soi.product_name, soi.stock_quantity,
                   soi.red_list, soi.sea_freight, soi.notes, soi.sku
            FROM shipment_order_items soi
            WHERE soi.shipment_order_id = :oid AND soi.deleted_at IS NULL
        """), {"oid": order_id}).fetchall()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "发货单明细"
        ws.append(["产品编码", "SKU", "产品名称", "数量", "红单", "海运", "备注"])
        for it in items:
            ws.append([
                it[1] or "",
                it[7] or "",
                it[2] or "",
                int(it[3]) if it[3] else 0,
                it[4] or "",
                it[5] or "",
                it[6] or "",
            ])

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        from fastapi.responses import StreamingResponse
        from urllib.parse import quote
        filename = f"发货单_{order[1]}.xlsx"
        encoded_filename = quote(filename)
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"}
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"导出发货单失败: {str(e)}")


@router.post("/{order_id}/import")
async def import_shipment_detail(
    order_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(PermissionChecker("shipment:edit"))
):
    """导入Excel更新发货单明细的红单/海运数量"""
    try:
        order = db.execute(text("""
            SELECT id, status FROM shipment_orders
            WHERE id = :id AND tenant_id = :tid AND deleted_at IS NULL
        """), {"id": order_id, "tid": current_user.tenant_id}).fetchone()

        if not order:
            raise HTTPException(status_code=404, detail="发货单不存在")
        if order[1] != 'draft':
            raise HTTPException(status_code=400, detail="只有草稿状态的发货单才能导入")

        content = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        ws = wb.active

        rows = list(ws.iter_rows(min_row=2, values_only=True))
        updated_count = 0
        not_found_codes = []

        for row in rows:
            if not row or not row[0]:
                continue
            product_code = str(row[0]).strip()
            red_list = str(row[4]).strip() if row[4] else ""
            sea_freight = str(row[5]).strip() if row[5] else ""
            notes = str(row[6]).strip() if len(row) > 6 and row[6] else ""

            result = db.execute(text("""
                UPDATE shipment_order_items
                SET red_list = :rl, sea_freight = :sf, notes = :nt
                WHERE shipment_order_id = :oid AND product_code = :pc AND deleted_at IS NULL
            """), {"rl": red_list, "sf": sea_freight, "nt": notes, "oid": order_id, "pc": product_code})

            if result.rowcount > 0:
                updated_count += 1
            else:
                not_found_codes.append(product_code)

        db.commit()
        msg = f"成功更新 {updated_count} 条记录"
        if not_found_codes:
            msg += f"，{len(not_found_codes)} 条产品编码未匹配: {', '.join(not_found_codes[:10])}"

        return {"success": True, "message": msg, "updated_count": updated_count, "not_found": not_found_codes}
    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"导入发货单失败: {str(e)}")